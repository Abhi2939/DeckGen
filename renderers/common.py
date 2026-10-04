from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.dml.color import RGBColor

from style import PALETTE, FONT_NAME


def rgb(c):
    return PALETTE[c] if isinstance(c, str) else c


def new_slide(prs, bg="bg_light"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(bg)
    return slide


def _style_run(run, size, bold, color, italic=False):
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)


def add_text(slide, x, y, w, h, text, size=12, bold=False, color="neutral_dark",
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        _style_run(p.add_run(), size, bold, color, italic)
        p.runs[0].text = line
    return tb


def add_runs(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """runs = [(text, size, bold, color), ...] in a single wrapped paragraph."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    for text, size, bold, color in runs:
        r = p.add_run()
        r.text = text
        _style_run(r, size, bold, color)
    return tb


def add_box(slide, x, y, w, h, fill, kind=MSO_SHAPE.RECTANGLE, line=None, radius=None, line_w=0.75):
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(line_w)
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    sh.shadow.inherit = False
    return sh


def add_line(slide, x1, y1, x2, y2, color="border", width=1.0, dashed=False):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = rgb(color)
    ln.line.width = Pt(width)
    if dashed:
        ln.line.dash_style = MSO_LINE.DASH
    return ln


def add_title(slide, text, x=0.45, y=0.25, w=8.2, size=26, align=PP_ALIGN.LEFT):
    return add_text(slide, x, y, w, 0.55, text, size, True, "neutral_dark", align)


def add_brand_ring(slide, color="blue_royal"):
    ring = add_box(slide, 9.05, 0.22, 0.5, 0.5, color, kind=MSO_SHAPE.DONUT)
    ring.adjustments[0] = 0.18
    return ring
import os
import re
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

from renderers.common import new_slide, add_text, add_box, add_line, add_title, add_brand_ring

LOGO_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "assets", "logos")
COLS, CELL_W, CELL_H = 3, 1.38, 0.46


def _slug(name):
    return re.sub(r"[^a-z0-9]+", "", name.lower())


def _place_name(slide, name, x, y):
    """Logo from assets/logos/<slug>.png if it exists, otherwise a text chip."""
    path = os.path.join(LOGO_DIR, _slug(name) + ".png")
    if os.path.exists(path):
        pic = slide.shapes.add_picture(path, Inches(x + 0.05), Inches(y + 0.04), height=Inches(CELL_H - 0.08))
        max_w = Inches(CELL_W - 0.1)
        if pic.width > max_w:
            ratio = max_w / pic.width
            pic.width, pic.height = int(pic.width * ratio), int(pic.height * ratio)
        pic.left = int(Inches(x) + (Inches(CELL_W) - pic.width) / 2)
        pic.top = int(Inches(y) + (Inches(CELL_H) - pic.height) / 2)
    else:
        add_box(slide, x + 0.04, y + 0.04, CELL_W - 0.08, CELL_H - 0.08, "white",
                MSO_SHAPE.ROUNDED_RECTANGLE, line="border", radius=0.25)
        add_text(slide, x + 0.08, y + 0.04, CELL_W - 0.16, CELL_H - 0.08, name, 9.5, True, "navy",
                 PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def render(prs, facts):
    groups = facts.competitors[:4]
    slide = new_slide(prs, "white")
    add_title(slide, "Competitive Landscape - Global", x=0.45, w=8.4, size=26, align=PP_ALIGN.CENTER)
    add_brand_ring(slide)

    add_line(slide, 5.0, 1.15, 5.0, 5.4, "border", 1, dashed=True)
    add_line(slide, 0.4, 3.25, 9.6, 3.25, "border", 1, dashed=True)

    quad_w = 4.3
    for i, g in enumerate(groups):
        right, bottom = i % 2 == 1, i >= 2
        qx = 5.25 if right else 0.45
        pill_w = 3.2
        pill_x = qx + (quad_w - pill_w) / 2
        pill_y = 4.95 if bottom else 1.2
        add_box(slide, pill_x, pill_y, pill_w, 0.4, "green", MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        add_text(slide, pill_x, pill_y, pill_w, 0.4, g.category, 11.5, True, "white", PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

        names = g.names[:COLS * 3]
        rows = (len(names) + COLS - 1) // COLS
        grid_x = qx + (quad_w - COLS * CELL_W) / 2
        area_top, area_bot = (3.4, 4.85) if bottom else (1.7, 3.15)
        grid_y = area_top + ((area_bot - area_top) - rows * CELL_H) / 2
        for k, name in enumerate(names):
            _place_name(slide, name, grid_x + (k % COLS) * CELL_W, grid_y + (k // COLS) * CELL_H)
    return slide
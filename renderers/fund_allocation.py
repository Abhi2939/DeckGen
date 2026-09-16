from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

from style import PALETTE, FONT_NAME, TITLE_SIZE, BODY_SIZE, LABEL_SIZE


def render(prs, facts):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.5), Inches(8.8), Inches(0.8))
    tf = title_box.text_frame
    tf.text = "Ask and allocation"
    tf.paragraphs[0].runs[0].font.size = TITLE_SIZE
    tf.paragraphs[0].runs[0].font.name = FONT_NAME
    tf.paragraphs[0].runs[0].font.color.rgb = PALETTE["neutral_dark"]
    tf.paragraphs[0].runs[0].font.bold = True

    if facts.ask_amount:
        ask_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.8), Inches(2.6), Inches(1.0))
        tf = ask_box.text_frame
        tf.text = facts.ask_amount
        run = tf.paragraphs[0].runs[0]
        run.font.size = Pt(40)
        run.font.bold = True
        run.font.name = FONT_NAME
        run.font.color.rgb = PALETTE["primary"]

        label_box = slide.shapes.add_textbox(Inches(0.6), Inches(2.7), Inches(2.6), Inches(0.5))
        tf = label_box.text_frame
        tf.text = "Total ask"
        run = tf.paragraphs[0].runs[0]
        run.font.size = BODY_SIZE
        run.font.name = FONT_NAME
        run.font.color.rgb = PALETTE["neutral_dark"]

    chart_data = CategoryChartData()
    chart_data.categories = [item.label for item in facts.allocation]
    chart_data.add_series("Allocation", [item.percent for item in facts.allocation])

    x, y, cx, cy = Inches(3.6), Inches(1.5), Inches(5.8), Inches(4.5)
    graphic_frame = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, x, y, cx, cy, chart_data)
    chart = graphic_frame.chart

    chart.has_title = False
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.RIGHT
    chart.legend.include_in_layout = False
    chart.legend.font.size = LABEL_SIZE
    chart.legend.font.name = FONT_NAME

    plot = chart.plots[0]
    plot.has_data_labels = True
    data_labels = plot.data_labels
    data_labels.show_percentage = True
    data_labels.show_value = False
    data_labels.show_category_name = False
    data_labels.font.size = LABEL_SIZE
    data_labels.font.name = FONT_NAME
    data_labels.font.bold = True
    data_labels.font.color.rgb = PALETTE["white"]

    series = plot.series[0]
    for i, point in enumerate(series.points):
        color = PALETTE["chart_series"][i % len(PALETTE["chart_series"])]
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = RGBColor.from_string(color)

    return slide
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

from style import PALETTE, FONT_NAME
from renderers.common import (new_slide, add_text, add_runs, add_box, add_line,
                              add_title, add_brand_ring, rgb)

TERM_COLOR = {"short": "term_short", "medium": "term_medium", "long": "term_long"}


def _cell(cell, text, size=9, bold=False, color="neutral_dark", fill="white", align=PP_ALIGN.LEFT):
    cell.fill.solid()
    cell.fill.fore_color.rgb = rgb(fill)
    cell.margin_left = Inches(0.08)
    cell.margin_right = Inches(0.05)
    cell.margin_top = Inches(0.025)
    cell.margin_bottom = Inches(0.025)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = cell.text_frame.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name, r.font.size, r.font.bold = FONT_NAME, Pt(size), bold
    r.font.color.rgb = rgb(color)


def render(prs, facts):
    m = facts.market
    slide = new_slide(prs, "bg_light")
    add_title(slide, m.title)
    add_brand_ring(slide)
    if m.cagr:
        add_text(slide, 0.45, 0.78, 5, 0.3, f"CAGR: {m.cagr}", 14, True, "blue_royal")

    top, h = 1.2, 1.9
    add_box(slide, 0.45, top, 3.5, h, "blue_royal", MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    if m.end:
        add_text(slide, 0.7, top + 0.2, 2.0, 0.3, "Projected", 10, False, "accent_light")
        add_text(slide, 3.0, top + 0.18, 0.8, 0.3, str(m.end.year), 13, True, "white", PP_ALIGN.RIGHT)
        add_text(slide, 0.7, top + 0.5, 3.1, 0.6, m.end.value, 24, True, "white")
    if m.start:
        add_box(slide, 0.45, top + 0.95, 2.6, 0.95, "white", MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
        add_text(slide, 2.2, top + 1.02, 0.75, 0.3, str(m.start.year), 12, True, "neutral_dark", PP_ALIGN.RIGHT)
        add_text(slide, 0.65, top + 1.3, 2.35, 0.5, m.start.value, 18, True, "neutral_dark", anchor=MSO_ANCHOR.MIDDLE)

    add_box(slide, 4.15, top, 5.4, h, "white", MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)
    segs = (m.segments or [])[:4]
    if segs:
        col_w = 5.4 / len(segs)
        for i, s in enumerate(segs):
            x = 4.15 + i * col_w + 0.15
            w = col_w - 0.3
            if i:
                add_line(slide, 4.15 + i * col_w, top + 0.2, 4.15 + i * col_w, top + h - 0.2, "border", 0.75)
            add_text(slide, x, top + 0.15, w, 0.25, s.category, 10, True, "gray_text")
            add_runs(slide, x, top + 0.48, w, 0.55,
                     [(s.leader_text + " ", 9.5, False, "neutral_dark"), (s.leader_value, 13, True, "blue_royal")])
            arrow = add_box(slide, x, top + 1.03, 0.2, 0.22, "green", MSO_SHAPE.DOWN_ARROW)
            add_text(slide, x, top + 1.3, w, 0.35, s.fast_text, 9.5, False, "neutral_dark")
            add_text(slide, x, top + 1.58, w, 0.28, s.fast_cagr, 13, True, "blue_royal")

    drivers = (m.drivers or [])[:6]
    if drivers:
        col_w = [3.0, 1.9, 2.5, 1.8]
        row_h = 0.29
        gf = slide.shapes.add_table(len(drivers) + 1, 4, Inches(0.45), Inches(3.3),
                                    Inches(sum(col_w)), Inches(row_h * (len(drivers) + 1)))
        tbl = gf.table
        tbl.horz_banding = False
        for i, w in enumerate(col_w):
            tbl.columns[i].width = Inches(w)
        for r in range(len(drivers) + 1):
            tbl.rows[r].height = Inches(row_h)
        for c, head in enumerate(["Driver", "(~) % Impact on CAGR", "Geographic Relevance", "Impact Timeline"]):
            _cell(tbl.cell(0, c), head, 9, True, "white", "green")
        for r, d in enumerate(drivers, start=1):
            band = "white" if r % 2 else "row_alt"
            _cell(tbl.cell(r, 0), d.driver, fill=band)
            _cell(tbl.cell(r, 1), d.impact, fill=band)
            _cell(tbl.cell(r, 2), d.relevance, fill=band)
            _cell(tbl.cell(r, 3), d.timeline, fill=TERM_COLOR.get(d.term, "term_medium"))
    return slide
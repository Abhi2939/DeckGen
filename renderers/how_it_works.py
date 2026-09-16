from pptx.util import Inches, Pt, Emu
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

from style import PALETTE, FONT_NAME, TITLE_SIZE, BODY_SIZE, LABEL_SIZE


def render(prs, facts):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.5), Inches(8.8), Inches(0.8))
    tf = title_box.text_frame
    tf.text = "How it works"
    tf.paragraphs[0].runs[0].font.size = TITLE_SIZE
    tf.paragraphs[0].runs[0].font.name = FONT_NAME
    tf.paragraphs[0].runs[0].font.color.rgb = PALETTE["neutral_dark"]
    tf.paragraphs[0].runs[0].font.bold = True

    steps = facts.process_steps
    n = len(steps)

    slide_width = Inches(10)
    margin = Inches(0.6)
    usable_width = slide_width - 2 * margin
    circle_d = Inches(0.55)
    circle_y = Inches(2.0)
    spacing = usable_width // n if n else usable_width

    centers = [margin + spacing * i + spacing // 2 for i in range(n)]

    if n > 1:
        line = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            centers[0], circle_y + circle_d // 2,
            centers[-1], circle_y + circle_d // 2,
        )
        line.line.color.rgb = PALETTE["accent_light"]
        line.line.width = Pt(2)

    for i, (step, cx) in enumerate(zip(steps, centers)):
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, cx - circle_d // 2, circle_y, circle_d, circle_d
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = PALETTE["accent"]
        circle.line.fill.background()
        tf = circle.text_frame
        tf.text = str(i + 1)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.runs[0].font.size = Pt(16)
        p.runs[0].font.bold = True
        p.runs[0].font.name = FONT_NAME
        p.runs[0].font.color.rgb = PALETTE["white"]

        text_box = slide.shapes.add_textbox(
            cx - Inches(1.0), circle_y + circle_d + Inches(0.15), Inches(2.0), Inches(1.6)
        )
        tf = text_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = step.title
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.name = FONT_NAME
        run.font.color.rgb = PALETTE["neutral_dark"]

        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        run2 = p2.add_run()
        run2.text = step.description
        run2.font.size = LABEL_SIZE
        run2.font.name = FONT_NAME
        run2.font.color.rgb = PALETTE["neutral_dark"]

    return slide